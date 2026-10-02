import ast
import io
import re
import zipfile
from copy import copy
from datetime import date
from pathlib import Path
from types import SimpleNamespace
from typing import Optional

import openpyxl
import pandas as pd
import pytest
from openpyxl.utils import get_column_letter

from app.data.time_data import normalize_job_area_value
from app.exports.google_templates import load_template_sheet_workbook, workbook_to_bytes
from app.exports.timeentries_export import build_daily_import_rate_cells
from app.features.project_assignments import build_job_number_key, is_truthy, resolve_employees_for_jobs


def export_namespace(time_data, assignments, template_bytes, errors):
    """Run the page's export functions without Streamlit page initialization."""
    source_path = Path(__file__).resolve().parents[1] / "pages" / "10_Timesheet_Entry.py"
    tree = ast.parse(source_path.read_text(encoding="utf-8"))
    names = {
        "_clean_text_value", "_get_employee_list_value", "_get_employee_truck",
        "_get_employee_post_to_payroll", "_get_employee_night_shift",
        "_get_employee_daily_import", "_employee_info_from_row", "_employee_info_lookup",
        "_lookup_employee_details", "_normalize_employee_key", "_find_col",
        "_has_daily_import_column", "_is_blank_value", "_is_truthy",
        "_prepare_employee_entries", "_write_employee_to_daily_time",
        "_daily_time_employee_sort_key", "_daily_import_night_trade_class",
        "apply_daily_import_data_row_style", "create_template_exports",
    }
    definitions = [node for node in ast.walk(tree)
                   if isinstance(node, ast.FunctionDef) and node.name in names]
    namespace = {
        "pd": pd, "Optional": Optional, "io": io, "zipfile": zipfile,
        "copy": copy, "re": re, "openpyxl": openpyxl,
        "get_column_letter": get_column_letter,
        "build_job_number_key": build_job_number_key,
        "resolve_employees_for_jobs": resolve_employees_for_jobs,
        "assignment_is_truthy": is_truthy,
        "build_daily_import_rate_cells": build_daily_import_rate_cells,
        "_normalize_job_area_value": normalize_job_area_value,
        "time_data_for_export": time_data, "XLSX": None,
        "st": SimpleNamespace(error=errors.append, warning=errors.append),
        "safe_read_excel": lambda _, sheet: assignments.copy() if sheet == "Employee Job Assignments" else pd.DataFrame(),
        "smart_read_data": lambda sheet, **kwargs: assignments.copy(),
        "get_google_template_workbook_bytes": lambda: template_bytes,
        "load_template_sheet_workbook": load_template_sheet_workbook,
        "workbook_to_bytes": workbook_to_bytes,
    }
    exec(compile(ast.Module(body=definitions, type_ignores=[]), str(source_path), "exec"), namespace)
    return namespace


@pytest.mark.parametrize("include_daily_time", [True, False])
def test_historical_package_includes_inactive_employees_and_saved_pay_rows(include_daily_time):
    employees = [
        ["10", "ALEX", "100", "EMPL", "EJ3", "Y", "Y", "225", "FALSE"],
        ["20", "ELBERT WIEBE", "100", "EMPL", "EJ2", "Y", "", "", "TRUE"],
        ["30", "BLAIR", "100", "EMPL", "EA2", "Y", "", "", "TRUE"],
        ["40", "DAILY TIME ONLY", "100", "EMPL", "EA1", "", "", "", "FALSE"],
        ["10", "ALEX", "200", "EMPL", "SUPER", "", "", "999", "TRUE"],
    ]
    assignments = pd.DataFrame(employees, columns=[
        "Person Number", "Employee Name", "Job Number", "Time Record Type",
        "Override Trade Class", "Daily Import", "Night Shift", "Subsistence Rate", "Active",
    ])
    time_data = pd.DataFrame([
        {
            "Job Number": "100", "Job Area": "008", "Date": "2026-09-03",
            "Name": name, "Employee Number": number, "Trade Class": trade,
            "RT Hours": 8, "OT Hours": 4, "Night Shift": night,
            "Subsistence Rate": subsistence, "Indirect": "FALSE", "Cost Code": "10-110-53",
        }
        for number, name, trade, night, subsistence in [
            ("10", "ALEX", "EJ3", "Y", "225"),
            ("20", "ELBERT WEIBE", "EJ2", "", ""),
            ("30", "BLAIR", "EA2", "", ""),
            ("40", "DAILY TIME ONLY", "EA1", "", ""),
            ("50", "REMOVED EMPLOYEE", "EJ1", "", ""),
        ]
    ])
    template = openpyxl.Workbook()
    template.active.title = "TimeEntries"
    if include_daily_time:
        template.create_sheet("Daily Time")
    errors = []
    namespace = export_namespace(time_data, assignments, workbook_to_bytes(template), errors)
    package = namespace["create_template_exports"](date(2026, 9, 3))
    assert package is not None, errors
    with zipfile.ZipFile(io.BytesIO(package)) as archive:
        wb = openpyxl.load_workbook(io.BytesIO(archive.read("09-03-2026 - 100 - Daily Import.xlsx")))
        rows = list(wb.active.iter_rows(min_row=4, max_col=14, values_only=True))
        assert {row[3] for row in rows} == {"ALEX", "ELBERT WEIBE", "BLAIR", "REMOVED EMPLOYEE"}
        alex = [row for row in rows if row[2] == "10"]
        assert [(row[1], row[4], row[9], row[10], row[11], row[12]) for row in alex] == [
            ("EMPL", "EJ3", "211", 8, "Y", "NS"),
            ("EMPL", "EJ3", "212", 4, "Y", "NS"),
            ("EMPL", "NS", "211", 12, "Y", "NS"),
            ("SUBS", "EJ3", "261", 1, None, "225"),
        ]
        assert all(row[7] == "008" for row in rows)
        if include_daily_time:
            daily = openpyxl.load_workbook(io.BytesIO(archive.read("09-03-2026 - Daily Time.xlsx")))
            names = {daily.active.cell(row=index, column=1).value for index in range(32, 37)}
            assert names == set(time_data["Name"])
            assert errors == []
        else:
            assert len(errors) == 1 and "Daily Time" in errors[0]
